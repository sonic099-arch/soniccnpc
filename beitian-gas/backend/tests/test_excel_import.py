from datetime import datetime, timezone, timedelta
from io import BytesIO
from zipfile import ZipFile
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, func
from app.core.config import get_settings
from app.data_sources.excel import ExcelDataSource
from app.models.entities import Well, WellMeasurement, DataImportJob
from app.services.imports import staged_path

HEADERS=['井号','时间','瞬时流量(m³/d)']

@pytest.fixture(autouse=True)
def isolated_staging(monkeypatch,tmp_path):
    monkeypatch.setattr(get_settings(),'import_storage_dir',tmp_path/'uploads')

def workbook(rows,headers=HEADERS):
    book=Workbook();sheet=book.active;sheet.title='测量数据';sheet.append(headers)
    for row in rows: sheet.append(row)
    data=BytesIO();book.save(data)
    return data.getvalue()

def upload(client,data):
    response=client.post('/api/v1/excel/preview',files={'file':('演示.xlsx',data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert response.status_code==201,response.text
    return response.json()

def options(job,**kwargs):
    preview=job['mapping']['preview']
    return dict(confirmed=True,mapping=preview['suggested_mapping'],units=preview['suggested_units'],**kwargs)

def confirm(client,job,**kwargs):
    return client.post(f"/api/v1/excel/jobs/{job['id']}/confirm",json=options(job,**kwargs))

def test_preview_requires_confirmation_before_writing(client,session):
    job=upload(client,workbook([['演示X','2026-10-01 08:00',2400]]))
    assert job['status']=='待确认' and job['total_rows']==1
    assert job['mapping']['preview']['suggested_units']['instantaneous_flow_m3_h']=='m³/d'
    assert session.scalar(select(func.count()).select_from(WellMeasurement))==0
    response=client.post(f"/api/v1/excel/jobs/{job['id']}/confirm",json={**options(job),'confirmed':False})
    assert response.status_code==422

def test_units_timezone_missing_values_and_well_creation(client,session):
    job=upload(client,workbook([['演示X',datetime(2026,10,1,8),2400]]))
    response=confirm(client,job,create_missing_wells=True,default_area='演示A')
    assert response.status_code==200,response.text
    assert response.json()['success_rows']==1
    row=session.scalar(select(WellMeasurement))
    assert row.instantaneous_flow_m3_h==100
    assert row.measured_at==datetime(2026,10,1,tzinfo=timezone.utc)
    assert row.tubing_pressure_mpa is None and '缺失:tubing_pressure_mpa' in row.quality_flags
    assert row.well.area=='演示A' and row.is_demo

def test_existing_well_association_uses_stable_id(client,session,records):
    create,_=records;well=create('wells')
    job=upload(client,workbook([['演示-001','2026-10-01T08:00:00+08:00',2400]]))
    assert confirm(client,job).json()['success_rows']==1
    assert str(session.scalar(select(WellMeasurement)).well_id)==well['id']

def test_unknown_well_reported_without_implicit_creation(client,session):
    job=upload(client,workbook([['未知井','2026-10-01',2400]]))
    result=confirm(client,job).json()
    assert result['failed_rows']==1 and '井号不存在' in result['errors'][0]['reason']
    assert session.scalar(select(func.count()).select_from(Well))==0

def test_repeated_upload_and_confirm_no_duplicates(client,session):
    data=workbook([['演示X','2026-10-01T00:00:00Z',2400],['演示X','2026-10-01T08:00:00+08:00',2400]])
    job=upload(client,data)
    first=confirm(client,job,create_missing_wells=True).json()
    assert (first['success_rows'],first['duplicate_rows'])==(1,1)
    again=confirm(client,job,create_missing_wells=True).json()
    assert again['success_rows']==1 and again['duplicate_rows']==1
    another=upload(client,data)
    result=confirm(client,another).json()
    assert result['success_rows']==0 and result['duplicate_rows']==2
    assert session.scalar(select(func.count()).select_from(WellMeasurement))==1

def test_conflicting_existing_values_reported_and_not_overwritten(client,session):
    first=upload(client,workbook([['演示X','2026-10-01',2400]]))
    confirm(client,first,create_missing_wells=True)
    second=upload(client,workbook([['演示X','2026-10-01',4800]]))
    result=confirm(client,second).json()
    assert result['failed_rows']==1 and result['duplicate_rows']==0
    assert '未覆盖' in result['errors'][0]['reason']
    assert session.scalar(select(WellMeasurement)).instantaneous_flow_m3_h==100

def test_errors_report_actual_excel_row_numbers_and_counts(client):
    data=workbook([
        ['演示X','2026-10-01',2400],
        [None,'2026-10-02',2400],
        ['演示X','无效时间',2400],
        ['演示X','2026-10-03','不是数字'],
        ['演示X','2026-10-04',None],
        [None,None,None],
        ['演示X','2026-10-05',2400],
    ])
    job=upload(client,data);result=confirm(client,job,create_missing_wells=True).json()
    assert (result['total_rows'],result['success_rows'],result['failed_rows'],result['duplicate_rows'])==(6,2,4,0)
    assert [e['row'] for e in result['errors']]==[3,4,5,6]
    report=client.get(result['report_url'])
    assert report.status_code==200
    book=load_workbook(BytesIO(report.content))
    assert book['错误行'].max_row==5
    assert book['错误行']['A2'].value==3
    book.close()
    assert client.get(f"/api/v1/excel/jobs/{job['id']}").json()['failed_rows']==4

@pytest.mark.parametrize('value',['NaN','inf',float('inf'),'=1+1',True])
def test_reject_nonfinite_formula_boolean(client,value):
    # openpyxl把float inf写为空，也会作为“所有参数为空”失败。
    job=upload(client,workbook([['演示X','2026-10-01',value]]))
    result=confirm(client,job,create_missing_wells=True).json()
    assert result['failed_rows']==1 and result['success_rows']==0

@pytest.mark.parametrize('field,unit,raw,expected',[
    ('油压(kPa)','kPa',6200,6.2),('油压(bar)','bar',62,6.2),
    ('累计流量(万m³)','万m³',2,20000),('气流温度(K)','K',273.15,0),
    ('产液量(m³/h)','m³/h',.5,12),('瞬时流量(万方/天)','万m³/d',2,20000/24),
])
def test_supported_unit_conversions(field,unit,raw,expected):
    source=ExcelDataSource(workbook([['演示X','2026-10-01',raw]],['井号','时间',field]))
    preview=source.preview()
    source.mapping=preview['suggested_mapping'];source.units=preview['suggested_units']
    source.validate_mapping(preview['headers'])
    _,row=next((n,r) for n,r in source.rows() if n!=1)
    record,_=source.parse_row(row)
    numeric_key=next(k for k in source.mapping if k not in ('well_number','measured_at'))
    assert source.units[numeric_key]==unit
    assert getattr(record,numeric_key)==pytest.approx(expected)

def test_invalid_mapping_and_units_leave_job_pending(client):
    job=upload(client,workbook([['演示X','2026-10-01',2400]]))
    base=options(job)
    for change in [{'mapping':{}},{'units':{}},{'timezone':'not-a-zone'},{'units':{'instantaneous_flow_m3_h':'MPa'}}]:
        response=client.post(f"/api/v1/excel/jobs/{job['id']}/confirm",json={**base,**change})
        assert response.status_code==422
    assert client.get(f"/api/v1/excel/jobs/{job['id']}").json()['status']=='待确认'

def test_invalid_extension_and_corrupt_files(client):
    for name,data in [('演示.xls',b'not excel'),('演示.xlsx',b'not excel')]:
        assert client.post('/api/v1/excel/preview',files={'file':(name,data)}).status_code==422

def test_duplicate_headers_rejected(client):
    assert client.post('/api/v1/excel/preview',files={'file':('x.xlsx',workbook([['x','y',1]],['井号','时间','井号']))}).status_code==422

def test_ambiguous_aliases_require_manual_mapping():
    preview=ExcelDataSource(workbook([['a','b','2026-10-01',1]],['井号','井名','时间','流量'])).preview()
    assert preview['ambiguous_fields']==['well_number'] and 'well_number' not in preview['suggested_mapping']

def test_tampered_staging_file_rejected(client):
    from uuid import UUID
    job=upload(client,workbook([['演示X','2026-10-01',2400]]))
    staged_path(UUID(job['id'])).write_bytes(b'changed')
    response=confirm(client,job,create_missing_wells=True)
    assert response.status_code==422 and '校验失败' in response.json()['detail']

def test_excel_datasource_contract_and_timezone_choice():
    source=ExcelDataSource(workbook([['演示X','2026-10-01 08:00',2400]]),timezone_name='UTC')
    preview=source.preview();source.mapping=preview['suggested_mapping'];source.units=preview['suggested_units']
    start=datetime(2026,10,1,8,tzinfo=timezone.utc)
    records=list(source.read_measurements(start,start+timedelta(hours=1)))
    assert records[0].measured_at==start
    assert list(source.read_measurements(start-timedelta(hours=1),start))==[]

def test_template_download_and_fields(client):
    fields=client.get('/api/v1/excel/fields').json()
    assert len(fields)==9
    response=client.get('/api/v1/excel/template?example=true')
    assert response.status_code==200
    source=ExcelDataSource(response.content);preview=source.preview()
    assert preview['total_rows']==3 and len(preview['suggested_mapping'])==9

def test_report_never_executes_user_formula(client):
    job=upload(client,workbook([['=HYPERLINK("http://example.invalid")','2026-10-01',2400]]))
    result=confirm(client,job,create_missing_wells=True).json()
    report=client.get(result['report_url'])
    book=load_workbook(BytesIO(report.content),data_only=False)
    assert book['错误行']['C2'].data_type=='s'
    book.close()

def test_chunked_import_over_five_hundred_rows(client,session):
    start=datetime(2026,10,1,tzinfo=timezone.utc)
    rows=[['演示X',(start+timedelta(hours=i)).isoformat(),2400] for i in range(505)]
    job=upload(client,workbook(rows))
    result=confirm(client,job,create_missing_wells=True).json()
    assert result['success_rows']==505 and result['failed_rows']==0
    second=upload(client,workbook(rows))
    assert confirm(client,second).json()['duplicate_rows']==505
    assert session.scalar(select(func.count()).select_from(WellMeasurement))==505

@pytest.mark.parametrize('timestamp',['now','today','2026-11-01 01:30:00','2026-03-08 02:30:00'])
def test_relative_or_ambiguous_times_never_imported(timestamp):
    source=ExcelDataSource(workbook([['演示X',timestamp,2400]]),timezone_name='America/New_York')
    preview=source.preview();source.mapping=preview['suggested_mapping'];source.units=preview['suggested_units']
    with pytest.raises(ValueError,match='时间'):
        source.parse_row({'井号':'演示X','时间':timestamp,'瞬时流量(m³/d)':2400})

def test_header_unit_suffix_without_parentheses():
    source=ExcelDataSource(workbook([['演示X','2026-10-01',6200]],['井号','时间','油压_kPa']))
    preview=source.preview()
    assert preview['suggested_mapping']['tubing_pressure_mpa']=='油压_kPa'
    assert preview['suggested_units']['tubing_pressure_mpa']=='kPa'

def test_blank_new_well_area_rejected(client):
    job=upload(client,workbook([['演示X','2026-10-01',2400]]))
    response=confirm(client,job,create_missing_wells=True,default_area='   ')
    assert response.status_code==422
