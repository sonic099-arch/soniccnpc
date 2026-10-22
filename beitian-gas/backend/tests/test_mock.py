from datetime import datetime, timezone, timedelta
from io import BytesIO
from uuid import UUID
import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from app.data_sources.mock import MockDataSource, SCENARIOS, PARAMETERS, DISCLAIMER
from app.services.demo import workbook_bytes, dataset_manifest, load_dataset, template_bytes
from app.models.entities import Well, WellMeasurement, WellMeasureEvent, WellOperationEvent, WellConfigHistory, WellTool, DataImportJob

@pytest.fixture(scope="module")
def small(): return MockDataSource(count=13,days=4).generate()

@pytest.fixture(scope="module")
def full(): return MockDataSource().generate()

def test_six_hundred_wells_and_thirteen_scenarios(full):
    assert len(full.wells)==600
    assert len({w['id'] for w in full.wells})==600
    assert {e['scenario'] for e in full.expected}=={s[0] for s in SCENARIOS}
    assert len(full.measurements)==201324
    assert all(w['is_demo'] and '演示' in w['description'] for w in full.wells)
    assert full.measurements['is_demo'].all()

def test_fixed_seed_reproducible(small):
    second=MockDataSource(count=13,days=4).generate()
    pd.testing.assert_frame_equal(small.measurements,second.measurements)
    assert small.wells==second.wells and small.measures==second.measures
    assert small.expected==second.expected
    assert workbook_bytes(small)==workbook_bytes(second)
    other=MockDataSource(count=13,days=4,seed=5).generate()
    assert not small.measurements.equals(other.measurements)

def scenario_frame(data,code):
    expected=next(e for e in data.expected if e['scenario']==code)
    return data.measurements.loc[data.measurements.well_number==expected['well_number']].reset_index(drop=True)

@pytest.mark.parametrize('code',[s[0] for s in SCENARIOS])
def test_scenario_expected_signal_and_label(small,code):
    frame=scenario_frame(small,code)
    expected=next(e for e in small.expected if e['scenario']==code)
    assert expected['is_demo'] and expected['warning']==DISCLAIMER
    assert all(k in expected for k in ['expected_production_state','expected_screening_status','expected_labels'])
    q=frame.instantaneous_flow_m3_h
    if code=='stable': assert q.max()/q.min()<1.1
    if code=='natural_decline': assert .85 < q.iloc[-1]/q.iloc[0] < .95
    if code in ('abnormal_decline','liquid_hint','freeze_hint','foam_poor'): assert q.iloc[-1]/q.iloc[0]<.6
    if code=='liquid_hint': assert (frame.casing_pressure_mpa-frame.tubing_pressure_mpa).iloc[-1]>2
    if code=='freeze_hint': assert frame.gas_temperature_c.iloc[-1]<0
    if code=='foam_normal': assert q.iloc[-1]/q.iloc[0]>.9 and expected['expected_screening_status']=='措施影响'
    if code=='cyclic': assert (q==0).sum()==64
    if code in ('planned_stop','unplanned_stop'): assert (q.tail(24)==0).all()
    if code=='frozen_data': assert frame[PARAMETERS].tail(24).nunique().max()==1
    if code=='missing_data': assert len(frame)==90 and q.isna().sum()==18 and expected['expected_screening_status']!='可信正常'
    if code=='counter_reset': assert frame.cumulative_flow_m3.diff().min()<0 and (q>0).all()
    if code in ('liquid_hint','freeze_hint'): assert any('辅助提示' in label for label in expected['expected_labels'])

def test_correlated_curves_and_cumulative_integration(small):
    frame=scenario_frame(small,'abnormal_decline')
    assert frame.instantaneous_flow_m3_h.corr(frame.gas_temperature_c)>.99
    assert frame.instantaneous_flow_m3_h.corr(frame.back_pressure_mpa)>.95
    q=frame.instantaneous_flow_m3_h.to_numpy()
    np.testing.assert_allclose(np.diff(frame.cumulative_flow_m3), (q[1:]+q[:-1])/2,atol=2e-6)

def test_event_timing_and_planned_stop(small):
    id_by_number={w['well_number']:w['id'] for w in small.wells}
    for code in ('planned_stop','unplanned_stop'):
        e=next(x for x in small.expected if x['scenario']==code)
        op=next(x for x in small.operations if x['well_id']==id_by_number[e['well_number']])
        assert op['planned']==(code=='planned_stop')
        assert op['effective_start'].isoformat()==e['onset']
    for event in small.configurations+small.tools+small.measures+small.operations:
        assert event['effective_start'].tzinfo is not None
        assert event['effective_end'] is None or event['effective_end']>event['effective_start']

def test_default_covers_methanol_and_other_measures(full):
    assert {'泡排','注醇','机械排液'}.issubset({e['measure_type'] for e in full.measures})

def test_datasource_half_open_interval(small):
    source=MockDataSource(count=13,days=4)
    start=datetime(2026,10,1,tzinfo=timezone.utc)
    records=list(source.read_measurements(start,start+timedelta(hours=1)))
    assert len(records)==13 and all(r.measured_at==start and r.is_demo for r in records)

def test_mock_load_full_relations_and_idempotency(session,small):
    first=load_dataset(session,small)
    assert first['measurement_count']==1242 and not first['already_loaded']
    second=load_dataset(session,small)
    assert second['already_loaded'] and first['job_id']==second['job_id']
    assert session.scalar(select(func.count()).select_from(Well))==13
    assert session.scalar(select(func.count()).select_from(WellMeasurement))==1242
    assert session.scalar(select(func.count()).select_from(WellConfigHistory))==13
    assert session.scalar(select(func.count()).select_from(WellTool))==13
    assert session.scalar(select(func.count()).select_from(WellMeasureEvent))==len(small.measures)
    assert session.scalar(select(func.count()).select_from(WellOperationEvent))==len(small.operations)
    with pytest.raises(ValueError,match='配置不同'):
        load_dataset(session,MockDataSource(count=13,days=4,seed=1).generate())

def test_demo_endpoint_requires_confirmation_and_is_idempotent(client):
    assert client.post('/api/v1/demo/generate',json={'confirmed':False}).status_code==422
    options={'confirmed':True,'count':13,'days':4}
    first=client.post('/api/v1/demo/generate',json=options)
    assert first.status_code==201,first.text
    assert not first.json()['already_loaded']
    assert client.post('/api/v1/demo/generate',json=options).json()['already_loaded']
    assert len(client.get('/api/v1/demo/scenarios').json()['scenarios'])==13

def test_template_and_workbook_sheets(small):
    book=load_workbook(BytesIO(workbook_bytes(small)),read_only=True)
    assert {'测量数据','气井信息','措施事件','操作事件','场景预期','说明'}.issubset(book.sheetnames)
    book.close()
    assert template_bytes()==template_bytes()

@pytest.mark.parametrize('options',[{'count':12},{'days':3},{'seed':-1},{'start':datetime(2026,1,1)}])
def test_reject_bad_generator_options(options):
    with pytest.raises(ValueError): MockDataSource(**options)

def test_loaded_events_relate_to_well_and_configuration(session,small):
    load_dataset(session,small)
    for expected in small.expected:
        well=session.scalar(select(Well).where(Well.well_number==expected['well_number']))
        config=session.scalar(select(WellConfigHistory).where(WellConfigHistory.well_id==well.id))
        assert config.configuration['demo_scenario']==expected['scenario']
        if expected['scenario'].startswith('foam'):
            event=session.scalar(select(WellMeasureEvent).where(WellMeasureEvent.well_id==well.id))
            assert event.measure_type=='泡排'
            assert event.effective_start.isoformat()==expected['onset']
