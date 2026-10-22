"""标准 xlsx 接入：识别只是建议，映射、单位、无时区时间的解释须显式确认。"""
from datetime import datetime, timezone
from io import BytesIO
import math
import re
from zoneinfo import ZoneInfo
from zipfile import BadZipFile, ZipFile
from openpyxl import load_workbook
import pandas as pd
from app.data_sources.base import MeasurementRecord

FIELDS = {
    "well_number": ("井号", None, ["井号", "井名", "气井编号", "well", "well_no", "well_number"]),
    "measured_at": ("采样时间", None, ["时间", "日期", "采样时间", "采集时间", "生产时间", "记录时间", "timestamp", "datetime", "measured_at"]),
    "tubing_pressure_mpa": ("油压", "MPa", ["油压", "油管压力", "tubing_pressure", "tubing_pressure_mpa"]),
    "casing_pressure_mpa": ("套压", "MPa", ["套压", "套管压力", "casing_pressure", "casing_pressure_mpa"]),
    "back_pressure_mpa": ("井口回压", "MPa", ["回压", "井口回压", "背压", "back_pressure", "back_pressure_mpa"]),
    "instantaneous_flow_m3_h": ("瞬时流量", "m³/h", ["瞬时流量", "瞬时产气量", "流量", "instantaneous_flow", "instantaneous_flow_m3_h"]),
    "cumulative_flow_m3": ("累计流量", "m³", ["累计流量", "累计产气量", "累积流量", "cumulative_flow", "cumulative_flow_m3"]),
    "gas_temperature_c": ("气流温度", "℃", ["温度", "气流温度", "井口温度", "gas_temperature", "gas_temperature_c"]),
    "liquid_rate_m3_d": ("产液量", "m³/d", ["产液量", "日产液量", "液量", "liquid_rate", "liquid_rate_m3_d"]),
}
UNITS = {
    "MPa": (1.0, 0.0), "kPa": (0.001, 0.0), "Pa": (0.000001, 0.0), "bar": (0.1, 0.0),
    "m³/h": (1.0, 0.0), "m³/d": (1/24, 0.0), "万m³/d": (10000/24, 0.0), "万m³/h": (10000.0, 0.0),
    "m³": (1.0, 0.0), "万m³": (10000.0, 0.0), "℃": (1.0, 0.0), "K": (1.0, -273.15),
}
ALLOWED_UNITS = {k: (["MPa", "kPa", "Pa", "bar"] if "pressure" in k else
                     ["m³/h", "m³/d", "万m³/d", "万m³/h"] if k == "instantaneous_flow_m3_h" else
                     ["m³", "万m³"] if k == "cumulative_flow_m3" else
                     ["℃", "K"] if k == "gas_temperature_c" else ["m³/d", "m³/h"])
                 for k, (_, u, _) in FIELDS.items() if u}
MAX_UPLOAD = 25 * 1024 * 1024
MAX_ROWS = 250_000
MAX_COLUMNS = 100

def normal(text):
    return re.sub(r"[\s_°]", "", str(text).lower().replace("³", "3").replace("℃", "c").replace("方", "m3").replace("天", "d").replace("小时", "h"))

def cell_json(value):
    if isinstance(value, datetime): return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value): return None
    if value is None or isinstance(value, (str, int, float, bool)): return value
    return str(value)

def inspect_archive(data):
    if len(data) > MAX_UPLOAD: raise ValueError("文件超过25 MiB，请拆分上传")
    try:
        with ZipFile(BytesIO(data)) as archive:
            info = archive.infolist()
            if len(info) > 2000 or sum(x.file_size for x in info) > 200*1024*1024:
                raise ValueError("Excel解压内容过大")
            if "xl/workbook.xml" not in archive.namelist(): raise ValueError("不是有效的xlsx工作簿")
    except BadZipFile as exc:
        raise ValueError("文件损坏或不是xlsx格式") from exc

class ExcelDataSource:
    def __init__(self, data: bytes, mapping=None, units=None, timezone_name="Asia/Shanghai", sheet=None, is_demo=True):
        self.data, self.mapping, self.units = data, mapping or {}, units or {}
        self.timezone_name, self.sheet, self.is_demo = timezone_name, sheet, is_demo

    def rows(self):
        inspect_archive(self.data)
        try:
            book = load_workbook(BytesIO(self.data), read_only=True, data_only=False, keep_links=False)
        except Exception as exc:
            raise ValueError("Excel无法读取，请检查工作簿格式") from exc
        try:
            name = self.sheet or ("测量数据" if "测量数据" in book.sheetnames else book.sheetnames[0])
            if name not in book.sheetnames: raise ValueError("所选工作表不存在")
            worksheet = book[name]
            if worksheet.max_column and worksheet.max_column > MAX_COLUMNS:
                raise ValueError("工作表超过100列")
            iterator = worksheet.iter_rows(values_only=True)
            first = next(iterator, ())
            while first and first[-1] is None: first = first[:-1]
            if not first: raise ValueError("首行必须是字段表头")
            headers = [str(v).strip() if v is not None else "" for v in first]
            if len(headers) > MAX_COLUMNS or any(not h for h in headers) or len(set(headers)) != len(headers):
                raise ValueError("表头不能重复、为空或超过100列")
            yield 1, {"headers": headers, "sheets": book.sheetnames, "sheet": name}
            count = 0
            for rowno, values in enumerate(iterator, 2):
                if len(values) > MAX_COLUMNS: raise ValueError("工作表超过100列")
                if not any(v is not None and str(v).strip() for v in values): continue
                count += 1
                if count > MAX_ROWS: raise ValueError("单次导入最多25万条记录")
                yield rowno, dict(zip(headers, values))
        finally:
            book.close()

    def preview(self):
        iterator = self.rows()
        _, meta = next(iterator)
        mapping, units, ambiguous = {}, {}, []
        for header in meta["headers"]:
            base = re.split(r"[（(\[]", header)[0]
            for key, (_, standard, aliases) in FIELDS.items():
                suffix_units = [u for u in ALLOWED_UNITS.get(key, []) if any(normal(base) == normal(a) + normal(u) for a in aliases)]
                if normal(base) in {normal(a) for a in aliases} or suffix_units:
                    if key in mapping:
                        ambiguous.append(key)
                    else:
                        mapping[key] = header
                        if standard:
                            unit_part = re.search(r"[（(\[]([^）)\]]+)", header)
                            unit_text = normal(unit_part[1]) if unit_part else normal(suffix_units[0]) if suffix_units else normal(standard)
                            matches = [u for u in ALLOWED_UNITS[key] if normal(u) == unit_text]
                            units[key] = matches[0] if matches else ""
        samples, total = [], 0
        for rowno, row in iterator:
            total += 1
            if len(samples) < 5: samples.append({"row": rowno, "values": {k: cell_json(v) for k,v in row.items()}})
        for key in set(ambiguous):
            mapping.pop(key, None); units.pop(key, None)
        return {**meta, "suggested_mapping": mapping, "suggested_units": units,
                "ambiguous_fields": sorted(set(ambiguous)), "total_rows": total, "samples": samples}

    def validate_mapping(self, headers):
        ZoneInfo(self.timezone_name)
        if not {"well_number", "measured_at"}.issubset(self.mapping): raise ValueError("必须映射井号和采样时间")
        if any(k not in FIELDS for k in self.mapping): raise ValueError("包含未知的标准字段")
        if any(v not in headers for v in self.mapping.values()): raise ValueError("映射引用了不存在的表头")
        if len(set(self.mapping.values())) != len(self.mapping): raise ValueError("同一个源列不能映射到多个参数")
        numeric = [k for k in self.mapping if k in ALLOWED_UNITS]
        if not numeric: raise ValueError("至少映射一个动态参数")
        for key in numeric:
            if self.units.get(key) not in ALLOWED_UNITS[key]: raise ValueError(f"{FIELDS[key][0]}的单位不受支持，请显式选择单位")

    def parse_row(self, row):
        well = row.get(self.mapping["well_number"])
        if well is None or isinstance(well, bool) or not str(well).strip(): raise ValueError("井号为空或无效")
        well = str(well).strip()
        if len(well) > 80 or well.startswith("="): raise ValueError("井号过长或包含公式")
        rawtime = row.get(self.mapping["measured_at"])
        if rawtime is None or isinstance(rawtime, (int, float, bool)): raise ValueError("时间为空或格式无效，请使用Excel日期或时间字符串")
        try:
            if isinstance(rawtime, str) and not re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", rawtime.strip()):
                raise ValueError("时间必须是固定日期，不支持now或today等相对时间")
            ts = pd.Timestamp(rawtime)
            if pd.isna(ts): raise ValueError("空时间")
            dt = ts.to_pydatetime()
            if dt.tzinfo is None:
                zone = ZoneInfo(self.timezone_name)
                if dt.replace(tzinfo=zone, fold=0).utcoffset() != dt.replace(tzinfo=zone, fold=1).utcoffset():
                    raise ValueError("夏令时边界时间不明确，请在源数据中指定时区偏移")
                dt = dt.replace(tzinfo=zone)
            dt = dt.astimezone(timezone.utc)
        except Exception as exc:
            raise ValueError("时间无效或存在时区歧义") from exc
        values, flags, present = {}, [], 0
        for key in ALLOWED_UNITS:
            raw = row.get(self.mapping[key]) if key in self.mapping else None
            if raw is None or (isinstance(raw, str) and not raw.strip()):
                values[key] = None
                flags.append("缺失:" + key)
                continue
            try:
                if isinstance(raw, bool) or (isinstance(raw, str) and raw.startswith("=")): raise ValueError()
                value = float(str(raw).strip())
                if not math.isfinite(value): raise ValueError()
                unit = self.units[key]
                scale, offset = UNITS[unit]
                if key == "liquid_rate_m3_d": scale = 24.0 if unit == "m³/h" else 1.0
                result = value*scale+offset
                if not math.isfinite(result): raise ValueError()
                values[key] = result
                present += 1
                if result < 0 and key != "gas_temperature_c": flags.append("负值待核实:" + key)
            except (ValueError, OverflowError) as exc:
                raise ValueError(f"{FIELDS[key][0]}不是有限数值（公式不允许导入）") from exc
        if not present: raise ValueError("所有动态参数均为空")
        record = MeasurementRecord(well_number=well, measured_at=dt, is_demo=self.is_demo, **values)
        return record, flags

    def read_measurements(self, start, end):
        iterator = self.rows(); _, meta = next(iterator)
        self.validate_mapping(meta["headers"])
        for rowno, row in iterator:
            try:
                record, _ = self.parse_row(row)
            except ValueError as exc:
                raise ValueError(f"第{rowno}行：{exc}") from exc
            if start <= record.measured_at < end: yield record
