"""模拟数据导出和显式入库；与分析引擎严格分离。"""
from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import json
import re
from uuid import NAMESPACE_URL, uuid5
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select, text, func, or_
from sqlalchemy.dialects.postgresql import insert
from app.data_sources.excel import FIELDS
from app.data_sources.mock import DISCLAIMER, MockDataset, MockDataSource
from app.models.entities import DataImportJob, Well, WellMeasurement, WellConfigHistory, WellTool, WellMeasureEvent, WellOperationEvent

HEADERS = ["井号", "时间", "油压(MPa)", "套压(MPa)", "井口回压(MPa)", "瞬时流量(m³/h)", "累计流量(m³)", "气流温度(℃)", "产液量(m³/d)"]

def deterministic_xlsx(book):
    # 固定元数据与ZIP条目时间，确保同一模拟数据的文件校验和也可重复。
    fixed = datetime(2026, 1, 1)
    book.properties.created = fixed; book.properties.modified = fixed
    stream = BytesIO(); book.save(stream)
    result = BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source, ZipFile(result, "w", ZIP_DEFLATED) as dest:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>).*?(</dcterms:modified>)", rb"\g<1>2026-01-01T00:00:00Z\g<2>", data)
            info = ZipInfo(name, (2026,1,1,0,0,0)); info.compress_type = ZIP_DEFLATED
            dest.writestr(info, data)
    return result.getvalue()


def style(book):
    for sheet in book:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17776D")
        for index in range(1, min(sheet.max_column,15)+1):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(index)].width = 25


def template_bytes(example=False):
    book = Workbook(); sheet = book.active; sheet.title = "测量数据"
    sheet.append(HEADERS)
    if example:
        sheet.append(["演示-导入-001", "2026-10-01T08:00:00+08:00",6.2,7.0,2.1,800,19200,18.5,0.8])
        sheet.append(["演示-导入-001", "2026-10-01T09:00:00+08:00",6.21,7.01,2.1,800,20000,18.6,0.8])
        sheet.append(["演示-导入-002", "2026-10-01T08:00:00+08:00",6.4,7.2,2.2,600,14400,19.0,0.6])
    notes = book.create_sheet("说明")
    notes.append(["项目", "说明"])
    notes.append(["数据性质",DISCLAIMER])
    notes.append(["使用流程","上传后确认映射、单位和无时区时间的解释，再执行导入。"])
    notes.append(["井号","已有气井用井号关联稳定UUID；不存在时须显式允许建井并指定工区。"])
    notes.append(["时间","推荐ISO8601带时区文本；Excel日期按用户确认的时区解释。"])
    notes.append(["标准单位","压力MPa、瞬时流量m³/h、累计流量m³、温度℃、产液量m³/d。体积标况和压力口径仍需来源说明。"])
    notes.append(["空值","部分参数可为空，保存缺失标记；井号、时间或所有动态参数为空的行失败。"])
    notes.append(["重复","同井同时间同数值跳过；不同数值报告冲突，绝不覆盖历史。"])
    notes.append(["第七参数","暂定产液量，待现场口径确认；当前为演示字段。"])
    style(book)
    return deterministic_xlsx(book)


def dataset_manifest(dataset):
    return dict(schema_version="mock-v1", seed=dataset.seed, is_demo=True, warning=DISCLAIMER,
                well_count=len(dataset.wells), measurement_count=len(dataset.measurements),
                wells=dataset.wells, configurations=dataset.configurations, tools=dataset.tools,
                measures=dataset.measures, operations=dataset.operations, expected=dataset.expected)


def workbook_bytes(dataset):
    book = Workbook(write_only=True)
    def header(sheet, values):
        sheet.freeze_panes = "A2"
        cells=[]
        for value in values:
            cell=WriteOnlyCell(sheet,value=value)
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="17776D")
            cells.append(cell)
        sheet.append(cells)
    sheet = book.create_sheet("测量数据")
    header(sheet, HEADERS)
    frame = dataset.measurements.drop(columns=["is_demo"]).astype(object).where(pd.notna(dataset.measurements.drop(columns=["is_demo"])), None)
    for row in frame.itertuples(index=False, name=None):
        sheet.append([v.isoformat() if isinstance(v, datetime) else v for v in row])
    manifest = dataset_manifest(dataset)
    for title, key in [("气井信息","wells"),("生产配置","configurations"),("井筒工具","tools"),("措施事件","measures"),("操作事件","operations"),("场景预期","expected")]:
        ws = book.create_sheet(title)
        records = manifest[key]
        if not records:
            header(ws,["说明"]); ws.append(["本批次无记录"]); continue
        fields = list(records[0])
        header(ws, fields)
        for record in records:
            ws.append([json.dumps(record.get(k),ensure_ascii=False,default=str) if isinstance(record.get(k),(dict,list)) else
                       str(record[k]) if record.get(k) is not None and not isinstance(record.get(k),(bool,float,int,str)) else record.get(k) for k in fields])
    notes = book.create_sheet("说明")
    notes.append(["项目","说明"]); notes.append(["数据性质",DISCLAIMER]); notes.append(["随机种子",dataset.seed])
    notes.append(["来源","MockDataSource mock-v1；仅用于测试，不含现场数据。"])
    notes.append(["导入范围","Excel界面只导入测量数据；气井、配置、工具及措施/操作由生成脚本 --load 或演示生成接口统一入库。"])
    notes.append(["预期标签","只作测试参照，不是分析引擎的结论。疑似积液与冻堵不确诊。"])
    return deterministic_xlsx(book)



def _signature(dataset, well_count):
    prefix = {w["well_number"] for w in dataset.wells[:well_count]}
    frame = dataset.measurements.loc[dataset.measurements.well_number.isin(prefix)]
    start = frame.measured_at.min().isoformat()
    return sha256(json.dumps(["mock-v1",dataset.seed,well_count,len(frame),start], ensure_ascii=False).encode()).hexdigest()


def _validate_existing(session, dataset, existing):
    """仅允许同一生成配置的完整历史前缀扩容；既有测量和人工复核不作改写。"""
    if not existing: return
    message = "演示井号已存在但生成配置不同、井号已变更或原数据不完整；请保持种子、天数、起始时间一致，不会覆盖已有数据"
    expected = {w["well_number"]:w for w in dataset.wells}
    for well in existing:
        if well.well_number not in expected or not well.is_demo or well.id != expected[well.well_number]["id"]:
            raise ValueError(message)
    existing_by_id = {w.id:w for w in existing}
    # 原版本只在mapping保存种子，通过实际来源任务成员数兼容旧600井批次。
    groups = session.execute(select(WellMeasurement.well_id,WellMeasurement.import_job_id,
                WellMeasurement.source_type,WellMeasurement.is_demo,func.count(),
                func.min(WellMeasurement.measured_at),func.max(WellMeasurement.measured_at))
                .where(WellMeasurement.well_id.in_(existing_by_id))
                .group_by(WellMeasurement.well_id,WellMeasurement.import_job_id,
                          WellMeasurement.source_type,WellMeasurement.is_demo)).all()
    if len(groups) != len(existing): raise ValueError(message)
    stats = dataset.measurements.groupby("well_number").measured_at.agg(["count","min","max"])
    job_ids = {g[1] for g in groups}
    if None in job_ids: raise ValueError(message)
    jobs = {j.id:j for j in session.scalars(select(DataImportJob).where(DataImportJob.id.in_(job_ids)))}
    for job_id in job_ids:
        job = jobs.get(job_id)
        if job is None or job.source_type != "mock" or job.status != "已完成": raise ValueError(message)
        if job.mapping.get("generator") != "mock-v1" or job.mapping.get("seed") != dataset.seed: raise ValueError(message)
        source_count = job.mapping.get("source_well_count")
        if source_count is None:
            source_count = session.scalar(select(func.count(func.distinct(WellMeasurement.well_id)))
                                          .where(WellMeasurement.import_job_id==job_id))
        if not isinstance(source_count,int) or not 13 <= source_count <= len(dataset.wells): raise ValueError(message)
        if job.file_sha256 != _signature(dataset,source_count): raise ValueError(message)
    for well_id, job_id, source_type, is_demo, count, first, last in groups:
        expected_stats = stats.loc[existing_by_id[well_id].well_number]
        if source_type != "mock" or not is_demo or count != expected_stats["count"]:
            raise ValueError(message)
        if first != expected_stats["min"] or last != expected_stats["max"]: raise ValueError(message)


def load_dataset(session, dataset):
    signature = _signature(dataset,len(dataset.wells))
    session.execute(text("SELECT pg_advisory_xact_lock(1862400202)"))
    numbers = [w["well_number"] for w in dataset.wells]
    existing = list(session.scalars(select(Well).where(or_(Well.well_number.in_(numbers),Well.id.in_([w["id"] for w in dataset.wells])))))
    _validate_existing(session,dataset,existing)
    previous = session.scalar(select(DataImportJob).where(DataImportJob.source_type=="mock",DataImportJob.file_sha256==signature,DataImportJob.status=="已完成"))
    if previous and len(existing) == len(dataset.wells):
        session.commit()
        return {"job_id":str(previous.id),"well_count":len(dataset.wells),"measurement_count":len(dataset.measurements),
                "already_loaded":True,"added_wells":0,"added_measurements":0}
    present_numbers = {w.well_number for w in existing}
    new_wells = [w for w in dataset.wells if w["well_number"] not in present_numbers]
    new_ids = {w["id"] for w in new_wells}
    new_numbers = {w["well_number"] for w in new_wells}
    frame = dataset.measurements.loc[dataset.measurements.well_number.isin(new_numbers)]
    new_count = len(frame)
    job = DataImportJob(source_type="mock",file_sha256=signature,status="已完成",filename="确定性演示数据",
                        mapping={"generator":"mock-v1","seed":dataset.seed,"is_demo":True,"warning":DISCLAIMER,
                                 "source_well_count":len(dataset.wells),"added_wells":len(new_wells),
                                 "existing_wells":len(existing)},
                        total_rows=len(dataset.measurements),success_rows=new_count,
                        failed_rows=0,duplicate_rows=len(dataset.measurements)-new_count,errors=[],completed_at=datetime.now(timezone.utc))
    session.add(job); session.flush()
    if new_wells: session.execute(insert(Well), new_wells)
    for model, records in [(WellConfigHistory,dataset.configurations),(WellTool,dataset.tools),(WellMeasureEvent,dataset.measures),(WellOperationEvent,dataset.operations)]:
        records = [record for record in records if record["well_id"] in new_ids]
        if records:
            payload = [{**record,"id":uuid5(NAMESPACE_URL,f"beitian/mock/{dataset.seed}/{model.__tablename__}/{record['well_id']}/{record['effective_start']}/{record.get('operation_type','')}")} for record in records]
            for offset in range(0,len(payload),500): session.execute(insert(model),payload[offset:offset+500])
    well_ids = {w["well_number"]:w["id"] for w in new_wells}
    rows = frame.astype(object).where(pd.notna(frame),None).to_dict("records")
    for offset in range(0,len(rows),1000):
        payload=[]
        for row in rows[offset:offset+1000]:
            data = dict(row); number=data.pop("well_number")
            flags=["缺失:"+key for key,value in data.items() if value is None]
            payload.append({**data,"well_id":well_ids[number],"quality_flags":flags,"source_type":"mock","import_job_id":job.id})
        session.execute(insert(WellMeasurement),payload,execution_options={"render_nulls":True})
    session.commit()
    return {"job_id":str(job.id),"well_count":len(dataset.wells),"measurement_count":len(dataset.measurements),
            "already_loaded":not new_wells,"added_wells":len(new_wells),"added_measurements":new_count}
